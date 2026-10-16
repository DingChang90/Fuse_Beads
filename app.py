import io
from collections import Counter
import requests
import streamlit as st
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

# 页面基本配置
st.set_page_config(
    page_title="拼豆 Excel 图纸生成器",
    page_icon="🎨",
    layout="centered"
)

st.title("🎨 拼豆 Excel 图纸自动生成器")
st.markdown("上传一张图片，即可一键转换为专业的拼豆图纸 Excel 文件！")

# ==========================================
# 侧边栏：参数配置（带默认值）
# ==========================================
st.sidebar.header("⚙️ 图纸参数配置")
grid_width = st.sidebar.slider("图纸宽度 (格子数)", min_value=20, max_value=200, value=100, step=10)
grid_height = st.sidebar.slider("图纸高度 (格子数)", min_value=20, max_value=200, value=100, step=10)
color_count = st.sidebar.slider("颜色数量 (色系上限)", min_value=8, max_value=64, value=48, step=4)


# ==========================================
# 核心处理函数
# ==========================================
def create_pixel_matrix(image_file, width: int, height: int, color_count: int):
    image = Image.open(image_file).convert("RGB")
    image = image.resize((width, height), Image.Resampling.LANCZOS)

    quantized = image.quantize(colors=color_count, method=Image.Quantize.MEDIANCUT)
    rgb_image = quantized.convert("RGB")
    palette = quantized.getpalette()

    used = Counter(quantized.getdata())
    palette_colors = {}

    for index in used:
        base = index * 3
        palette_colors[index] = tuple(palette[base:base + 3])

    ordered_colors = [index for index, _ in used.most_common()]
    code_map = {index: f"C{i + 1:02d}" for i, index in enumerate(ordered_colors)}

    return quantized, rgb_image, palette_colors, ordered_colors, code_map, used


def rgb_to_hex(rgb):
    return "%02X%02X%02X" % rgb


def generate_excel_bytes(quantized, rgb_image, palette_colors, ordered_colors, code_map, used, width, height):
    wb = Workbook()
    ws = wb.active
    ws.title = "拼豆图纸"

    ws["A1"] = f"拼豆图纸｜{width} × {height}｜{len(ordered_colors)}色"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "每个方格代表 1 颗拼豆；右侧为颜色编号和数量"
    ws["A2"].font = Font(size=10)

    thin = Side(style="hair", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in range(height):
        excel_row = 4 + row
        ws.row_dimensions[excel_row].height = 11
        for col in range(width):
            excel_col = 1 + col
            color_index = quantized.getpixel((col, row))
            rgb = palette_colors[color_index]
            cell = ws.cell(excel_row, excel_col)
            cell.fill = PatternFill("solid", fgColor=rgb_to_hex(rgb))
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(excel_col)].width = 1.65

    legend_col = width + 3
    headers = ["编号", "颜色", "数量", "RGB"]
    for i, title in enumerate(headers):
        cell = ws.cell(3, legend_col + i, title)
        cell.font = Font(bold=True)

    for row, color_index in enumerate(ordered_colors, start=4):
        rgb = palette_colors[color_index]
        ws.cell(row, legend_col, code_map[color_index])
        color_cell = ws.cell(row, legend_col + 1)
        color_cell.fill = PatternFill("solid", fgColor=rgb_to_hex(rgb))
        color_cell.border = border
        ws.cell(row, legend_col + 2, used[color_index])
        ws.cell(row, legend_col + 3, f"{rgb[0]},{rgb[1]},{rgb[2]}")

    widths = [9, 8, 9, 16]
    for i, width_value in enumerate(widths):
        ws.column_dimensions[get_column_letter(legend_col + i)].width = width_value

    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

    # Sheet 2：编号施工图
    ws2 = wb.create_sheet("编号图纸")
    ws2["A1"] = f"编号施工图｜{width} × {height}"
    ws2["A1"].font = Font(bold=True, size=14)

    for row in range(height):
        excel_row = 3 + row
        ws2.row_dimensions[excel_row].height = 13
        for col in range(width):
            color_index = quantized.getpixel((col, row))
            rgb = palette_colors[color_index]
            cell = ws2.cell(excel_row, col + 1, code_map[color_index])
            cell.fill = PatternFill("solid", fgColor=rgb_to_hex(rgb))
            cell.border = border
            cell.font = Font(size=5, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws2.column_dimensions[get_column_letter(col + 1)].width = 3.0

    ws2.freeze_panes = "A3"
    ws2.sheet_view.showGridLines = False

    # Sheet 3：颜色统计
    ws3 = wb.create_sheet("颜色统计")
    ws3["A1"] = "拼豆颜色统计"
    ws3["A1"].font = Font(bold=True, size=14)
    for col, title in enumerate(["编号", "颜色", "数量", "RGB"], start=1):
        ws3.cell(3, col, title).font = Font(bold=True)

    for row, color_index in enumerate(ordered_colors, start=4):
        rgb = palette_colors[color_index]
        ws3.cell(row, 1, code_map[color_index])
        color_cell = ws3.cell(row, 2)
        color_cell.fill = PatternFill("solid", fgColor=rgb_to_hex(rgb))
        color_cell.border = border
        ws3.cell(row, 3, used[color_index])
        ws3.cell(row, 4, f"{rgb[0]},{rgb[1]},{rgb[2]}")

    ws3.column_dimensions["A"].width = 10
    ws3.column_dimensions["B"].width = 10
    ws3.column_dimensions["C"].width = 10
    ws3.column_dimensions["D"].width = 16
    ws3.sheet_view.showGridLines = False

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ==========================================
# 企业微信应用消息核心逻辑
# ==========================================
def get_wechat_access_token():
    """获取企业微信应用的 access_token"""
    try:
        cfg = st.secrets["wechat"]
        url = f"https://qyapi.weixin.qq.com/cgi-bin/gettoken?corpid={cfg['corp_id']}&corpsecret={cfg['corp_secret']}"
        res = requests.get(url, timeout=5).json()
        if res.get("errcode") == 0:
            return res.get("access_token")
    except Exception as e:
        print(f"获取微信access_token失败: {e}")
    return None


def upload_wechat_media(access_token, file_bytes, file_name, media_type):
    """上传临时素材到企业微信，返回 media_id (media_type: image 或 file)"""
    try:
        url = f"https://qyapi.weixin.qq.com/cgi-bin/media/upload?access_token={access_token}&type={media_type}"
        files = {"media": (file_name, file_bytes.getvalue())}
        res = requests.post(url, files=files, timeout=15).json()
        if "media_id" in res:
            return res["media_id"]
    except Exception as e:
        print(f"上传微信素材失败: {e}")
    return None


def send_wechat_notification(file_name, image_bytes, excel_bytes):
    """通过企业微信应用发送文本通知、图片及文件（含“拼豆”）"""
    try:
        if "wechat" not in st.secrets:
            return
        
        cfg = st.secrets["wechat"]
        access_token = get_wechat_access_token()
        if not access_token:
            return

        agent_id = cfg["agent_id"]
        to_user = cfg.get("to_user", "@all")
        base_url = f"https://qyapi.weixin.qq.com/cgi-bin/message/send?access_token={access_token}"
        headers = {"Content-Type": "application/json"}

        # 1. 发送文字通知（含“拼豆”）
        text_data = {
            "touser": to_user,
            "msgtype": "text",
            "agentid": agent_id,
            "text": {"content": f"【拼豆通知】有用户刚刚成功生成了一份拼豆图纸！\n文件名：{file_name}"}
        }
        requests.post(base_url, json=text_data, headers=headers, timeout=5)

        # 2. 上传并发送用户原图
        image_media_id = upload_wechat_media(access_token, image_bytes, "user_photo.jpg", "image")
        if image_media_id:
            img_data = {
                "touser": to_user,
                "msgtype": "image",
                "agentid": agent_id,
                "image": {"media_id": image_media_id}
            }
            requests.post(base_url, json=img_data, headers=headers, timeout=5)

        # 3. 上传并发送生成的 Excel 图纸文件
        file_media_id = upload_wechat_media(access_token, excel_bytes, file_name, "file")
        if file_media_id:
            file_data = {
                "touser": to_user,
                "msgtype": "file",
                "agentid": agent_id,
                "file": {"media_id": file_media_id}
            }
            requests.post(base_url, json=file_data, headers=headers, timeout=5)

    except Exception as e:
        print(f"发送企业微信通知失败: {e}")


# ==========================================
# 主界面逻辑
# ==========================================
uploaded_file = st.file_uploader("请上传需要转换的照片", type=["jpg", "jpeg", "png", "webp"])

if uploaded_file is not None:
    col1, col2 = st.columns(2)
    with col1:
        st.subheader("原图预览")
        st.image(uploaded_file, use_container_width=True)

    with col2:
        st.subheader("生成状态")
        if st.button("🚀 开始生成拼豆 Excel 图纸", type="primary"):
            with st.spinner("正在生成并同步至企业微信，请稍候..."):
                try:
                    # 1. 像素化与量化
                    quantized, rgb_image, palette_colors, ordered_colors, code_map, used = create_pixel_matrix(
                        uploaded_file, grid_width, grid_height, color_count
                    )

                    # 2. 生成 Excel 字节流
                    excel_data = generate_excel_bytes(
                        quantized, rgb_image, palette_colors, ordered_colors, code_map, used, grid_width, grid_height
                    )

                    base_name = uploaded_file.name.rsplit(".", 1)[0]
                    output_filename = f"拼豆图纸_{base_name}.xlsx"

                    # 3. 备份字节用于微信推送
                    uploaded_file.seek(0)
                    image_bytes_copy = io.BytesIO(uploaded_file.read())
                    excel_bytes_copy = io.BytesIO(excel_data.getvalue())

                    # 4. 触发企业微信应用通知（发文字、发图、发文件）
                    send_wechat_notification(output_filename, image_bytes_copy, excel_bytes_copy)

                    # 5. 网页端下载按钮（无多余成功提示文字）
                    st.download_button(
                        label="📥 点击下载生成的拼豆 Excel 图纸",
                        data=excel_data,
                        file_name=output_filename,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                except Exception as e:
                    st.error(f"生成过程中出现错误：{e}")
